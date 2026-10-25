"""Post-build acceptance checks. Run: python tests/validate_project.py"""
from pathlib import Path
import json
import sqlite3
import sys

import pandas as pd
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
DELIVERABLES = ROOT.parent / "deliverables"
DB = ROOT / "indus_weave.db"


def main() -> None:
    with sqlite3.connect(DB) as conn:
        checks = {
            "all 12 raw source tables loaded": conn.execute(
                "SELECT COUNT(*) FROM sqlite_master WHERE type='table' AND name LIKE 'raw_%'"
            ).fetchone()[0] == 12,
            "raw employee source count": conn.execute(
                "SELECT COUNT(*) FROM raw_employees"
            ).fetchone()[0] == 1316,
            "clean employee dimension count": conn.execute(
                "SELECT COUNT(DISTINCT emp_id) FROM dim_employee"
            ).fetchone()[0] == 1266,
            "attendance fact count": conn.execute(
                "SELECT COUNT(*) FROM fact_attendance_daily"
            ).fetchone()[0] == 217584,
            "attendance employee-date uniqueness": conn.execute(
                "SELECT COUNT(*) FROM (SELECT emp_id,date_key FROM fact_attendance_daily "
                "GROUP BY emp_id,date_key HAVING COUNT(*)>1)"
            ).fetchone()[0] == 0,
            "attendance hours nonnegative": conn.execute(
                "SELECT COUNT(*) FROM fact_attendance_daily "
                "WHERE worked_hours<0 OR computed_overtime_hours<0"
            ).fetchone()[0] == 0,
            "payroll line count": conn.execute(
                "SELECT COUNT(*) FROM fact_payroll_line"
            ).fetchone()[0] == 53617,
            "all foreign keys resolve": len(conn.execute("PRAGMA foreign_key_check").fetchall()) == 0,
            "all DQ checks pass": conn.execute(
                "SELECT COUNT(*) FROM dq_results WHERE status<>'PASS'"
            ).fetchone()[0] == 0,
        }

    analytics_dir = ROOT / "results" / "analytics"
    query_results_path = analytics_dir / "query_test_results.csv"
    query_results = pd.read_csv(query_results_path)
    checks["all ten scored SQL questions pass"] = (
        len(query_results) == 10 and (query_results["status"] == "PASS").all()
    )
    query_files_execute = True
    with sqlite3.connect(DB) as conn:
        for query_name in query_results["query"]:
            sql_file = ROOT / "analytics" / f"{query_name}.sql"
            output_file = analytics_dir / f"{query_name}.csv"
            if not sql_file.exists() or not output_file.exists():
                query_files_execute = False
                break
            cursor = conn.execute(sql_file.read_text(encoding="utf-8"))
            expected_output = pd.read_csv(output_file)
            if (
                [description[0] for description in cursor.description or []]
                != expected_output.columns.tolist()
                or len(cursor.fetchall()) != len(expected_output)
            ):
                query_files_execute = False
                break
    checks["all ten checked-in SQL files execute to their CSV outputs"] = query_files_execute
    checks["ten SQL result CSVs exist"] = all(
        (analytics_dir / f"{query}.csv").exists()
        for query in query_results["query"].tolist()
    )
    paid_without_attendance = pd.read_csv(
        analytics_dir / "q07_paid_zero_attendance.csv"
    )
    paid_after_separation = pd.read_csv(
        analytics_dir / "q08_paid_after_separation.csv"
    )
    checks["zero-attendance report is based on positive earnings"] = (
        {"emp_id", "month", "paid_earnings", "present_days"}.issubset(
            paid_without_attendance.columns
        )
        and (paid_without_attendance["present_days"] == 0).all()
        and (paid_without_attendance["paid_earnings"] > 0).all()
    )
    checks["post-separation report contains earnings only"] = (
        {"emp_id", "month_start", "separation_date", "post_separation_earnings", "earning_line_count"}.issubset(
            paid_after_separation.columns
        )
        and (paid_after_separation["post_separation_earnings"] > 0).all()
        and not paid_after_separation.duplicated(["emp_id", "month_start"]).any()
    )
    with sqlite3.connect(DB) as conn:
        checks["overtime reconciliation preserves total hours"] = (
            conn.execute(
            """SELECT
                 abs((SELECT SUM(computed_hours) FROM staging_overtime_reconciliation) -
                     (SELECT SUM(computed_overtime_hours) FROM fact_attendance_daily)) < 0.01
                 AND abs((SELECT SUM(approved_hours) FROM staging_overtime_reconciliation) -
                         (SELECT SUM(approved_hours) FROM fact_overtime_approved)) < 0.01
                 AND abs((SELECT SUM(paid_hours) FROM staging_overtime_reconciliation) -
                         (SELECT SUM(units) FROM fact_payroll_line WHERE upper(pay_component)='OVERTIME')) < 0.01"""
            ).fetchone()[0] == 1
        )
    checks["supplemental missing-punch evidence exists"] = (
        (analytics_dir / "review_missing_punches.csv").exists()
        and (ROOT / "analytics" / "review_missing_punches.sql").exists()
    )
    checks["recruitment funnel bonus output exists"] = (
        (analytics_dir / "recruitment_funnel.csv").exists()
    )

    dashboard = ROOT / "web" / "dashboard.html"
    checks["dashboard is generated with embedded data"] = (
        dashboard.exists() and "__DATA__" not in dashboard.read_text(encoding="utf-8")
    )
    findings = ROOT / "presentation" / "cfo_findings.md"
    checks["CFO findings include query evidence and actions"] = (
        findings.exists()
        and "q05_overtime_reconciliation.csv" in findings.read_text(encoding="utf-8")
        and "q08_paid_after_separation.csv" in findings.read_text(encoding="utf-8")
        and "Action:" in findings.read_text(encoding="utf-8")
    )

    deck = ROOT / "presentation" / "PeopleOps_HR_Pipeline_Results.pptx"
    deck_ok = deck.exists()
    if deck_ok:
        presentation = Presentation(deck)
        checks["presentation has at most six slides"] = len(presentation.slides) <= 6
        checks["presentation has six slides"] = len(presentation.slides) == 6
    else:
        checks["presentation has six slides"] = False

    published_dashboard = DELIVERABLES / "PeopleOps_HR_Dashboard.html"
    published_deck = DELIVERABLES / "PeopleOps_HR_Pipeline_Results.pptx"
    checks["dashboard copied to submission deliverables"] = (
        published_dashboard.exists()
        and dashboard.exists()
        and published_dashboard.read_bytes() == dashboard.read_bytes()
    )
    checks["presentation copied to submission deliverables"] = (
        published_deck.exists()
        and deck.exists()
        and published_deck.read_bytes() == deck.read_bytes()
    )
    erd = ROOT / "warehouse" / "erd.svg"
    checks["rendered warehouse ER diagram exists"] = (
        erd.exists() and "<svg" in erd.read_text(encoding="utf-8")
    )
    checks["declared dependencies exist"] = (ROOT / "requirements.txt").exists()

    for name, passed in checks.items():
        print(f"{'PASS' if passed else 'FAIL'} - {name}")
    if not all(checks.values()):
        sys.exit(1)
    print(f"All {len(checks)} project acceptance checks passed.")


if __name__ == "__main__":
    main()
